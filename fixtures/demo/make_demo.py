"""Builds the demo Bureau's files (plan docs/plans/2026-10-solo-mvp.md, S-3), in French or English.

    python -m venv venv && venv/bin/pip install openpyxl python-pptx python-docx
    venv/bin/python fixtures/demo/make_demo.py            # French, in fixtures/demo/
    venv/bin/python fixtures/demo/make_demo.py --lang en  # English, in fixtures/demo/en/

Needs `cargo` (the workbook is baked through oel-xlsx: openpyxl writes formulas without their
results and charts without caches). `logo.png` is made once from `logo.svg` (the site's
favicon, same mark as `packages/emboss/src/icons/Logo.tsx`):

    magick -background none -density 3072 fixtures/demo/logo.svg -resize 1024x1024 -depth 8 fixtures/demo/logo.png

- « Chronos — tours de piste.xlsx »: lap times (sheet « Tours ») and results read from them
  with MINIFS, AVERAGEIFS, RANK.EQ, XLOOKUP and COUNTA (sheet « Résultats »), plus a chart.
- « Bienvenue dans le Bureau.pptx »: intro with the logo, three tips, the race results (a
  chart named « Graphique du tableur », replaced on the desk by a linked copy of the
  workbook's chart), « Merci ! » in white on black.
- « À propos du Bureau.docx »: the logo and a short text about the Bureau, with a shortcut table.

The English files (« Chronos — lap times.xlsx », « Welcome to the Bureau.pptx », « About the
Bureau.docx ») have the same content in English: the Bureau seeds a desk in the interface
language. Formulas are English in both. Fictional runners. Author « OutilsEnLigne », fixed
dates, so a run gives the same content.
"""
import argparse
import datetime as dt
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt as PPt

HERE = Path(__file__).parent
ROOT = HERE.parent.parent
LOGO = HERE / "logo.png"
AUTHOR = "OutilsEnLigne"
WHEN = dt.datetime(2026, 10, 11, 9, 0)

# Every word of the three files, per language (the Bureau's `src/bureau/demo.ts` names the same files, sheet and chart).
TEXT = {
    "fr": {
        "dir": HERE,
        "xlsx": "Chronos — tours de piste.xlsx",
        "pptx": "Bienvenue dans le Bureau.pptx",
        "docx": "À propos du Bureau.docx",
        "results": "Résultats",
        "laps": "Tours",
        "laps_header": ["Coureur", "Tour", "Temps (s)", "Date"],
        "results_header": ["Coureur", "Meilleur tour", "Moyenne", "Écart au meilleur", "Rang"],
        "summary": "En bref",
        "fastest": "Tour le plus rapide (s)",
        "by": "Par",
        "count": "Tours enregistrés",
        "group_mean": "Moyenne du groupe (s)",
        "chart_title": "Meilleur tour par coureur",
        "seconds": "Secondes",
        "deck_subtitle": "Tous vos fichiers sur un seul bureau",
        "how": "Comment l’utiliser",
        "tips": [
            "Double-cliquez sur un texte pour le modifier.",
            "Ajoutez une diapo avec « + » dans la pellicule.",
            "Changez un temps dans le tableur : le graphique de la diapo 3 suit.",
        ],
        "race": "Résultats de la course",
        "best": "Meilleur tour",
        "mean": "Moyenne",
        "deck_chart": "Graphique du tableur",
        "thanks": "Merci !",
        "thanks_name": "Merci",
        "deck_title": "Bienvenue dans le Bureau",
        "doc_title": "À propos du Bureau",
        "doc_text": (
            "Le Bureau d’OutilsEnLigne pose tous les fichiers d’un projet côte à côte, comme du papier sur une table : "
            "documents, présentations, tableurs, images et notes. Les outils suivent ce que vous touchez.",
            "Tout reste dans votre navigateur : vos fichiers ne quittent pas cet appareil. Pour en garder une copie, "
            "exportez le bureau depuis le menu, « Ouvrir et importer ».",
            "Chaque fichier se télécharge dans son format d’origine, Word, PowerPoint ou Excel ; ce que vous n’avez pas "
            "modifié y reste tel quel.",
        ),
        "shortcuts_title": "Raccourcis",
        "shortcuts_header": ("Pour", "Touches"),
        "shortcuts": [
            ("Ouvrir un fichier sur place", "Entrée ou double-clic"),
            ("Plein écran, et retour", "Ctrl + Maj + Entrée"),
            ("Revenir au bureau", "Échap"),
            ("Nouvelle note", "N"),
            ("Tous les raccourcis", "?"),
        ],
        "doc_end": "Ce document est un fichier Word ordinaire : il s’ouvre aussi dans Word ou LibreOffice.",
    },
    "en": {
        "dir": HERE / "en",
        "xlsx": "Chronos — lap times.xlsx",
        "pptx": "Welcome to the Bureau.pptx",
        "docx": "About the Bureau.docx",
        "results": "Results",
        "laps": "Laps",
        "laps_header": ["Runner", "Lap", "Time (s)", "Date"],
        "results_header": ["Runner", "Best lap", "Average", "Gap to the best", "Rank"],
        "summary": "In brief",
        "fastest": "Fastest lap (s)",
        "by": "By",
        "count": "Laps recorded",
        "group_mean": "Group average (s)",
        "chart_title": "Best lap per runner",
        "seconds": "Seconds",
        "deck_subtitle": "All your files on one desk",
        "how": "How to use it",
        "tips": [
            "Double-click a text to edit it.",
            "Add a slide with “+” in the filmstrip.",
            "Change a time in the spreadsheet: the chart on slide 3 follows.",
        ],
        "race": "Race results",
        "best": "Best lap",
        "mean": "Average",
        "deck_chart": "Spreadsheet chart",
        "thanks": "Thank you!",
        "thanks_name": "Thank you",
        "deck_title": "Welcome to the Bureau",
        "doc_title": "About the Bureau",
        "doc_text": (
            "The OutilsEnLigne Bureau lays every file of a project side by side, like paper on a table: "
            "documents, presentations, spreadsheets, images and notes. The tools follow what you touch.",
            "Everything stays in your browser: your files don’t leave this device. To keep a copy, export the "
            "desk from the menu, “Open and import”.",
            "Each file downloads in its original format, Word, PowerPoint or Excel; what you didn’t change stays "
            "exactly as it was.",
        ),
        "shortcuts_title": "Shortcuts",
        "shortcuts_header": ("To", "Keys"),
        "shortcuts": [
            ("Open a file in place", "Enter or double-click"),
            ("Full screen, and back", "Ctrl + Shift + Enter"),
            ("Back to the desk", "Esc"),
            ("New note", "N"),
            ("Every shortcut", "?"),
        ],
        "doc_end": "This document is an ordinary Word file: it also opens in Word or LibreOffice.",
    },
}
T = TEXT["fr"]

# 400 m, three laps each (seconds).
LAPS = {
    "Léa": (61.4, 60.2, 59.8),
    "Samuel": (66.1, 64.9, 65.3),
    "Maya": (58.7, 58.1, 59.0),
    "Noah": (70.2, 68.8, 69.5),
    "Zoé": (63.0, 62.4, 61.7),
    "Félix": (74.6, 72.9, 73.3),
}
LAP_DATES = (dt.date(2026, 10, 1), dt.date(2026, 10, 3), dt.date(2026, 10, 6))
RUNNERS = list(LAPS)
BEST = {r: min(t) for r, t in LAPS.items()}
MEAN = {r: round(sum(t) / len(t), 2) for r, t in LAPS.items()}

INK = "34333D"
VIOLET = "7859C5"
HEADER_FILL = PatternFill("solid", fgColor="EDE8F7")


# ── Workbook ────────────────────────────────────────────────────────────────────────────────


def header(ws, cells):
    for c in cells:
        c.font = Font(bold=True, color=INK)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")


def make_xlsx() -> None:
    wb = Workbook()
    res = wb.active
    res.title = T["results"]
    laps = wb.create_sheet(T["laps"])

    laps.append(T["laps_header"])
    header(laps, laps[1])
    for runner, times in LAPS.items():
        for i, t in enumerate(times):
            laps.append([runner, i + 1, t, LAP_DATES[i]])
    for row in laps.iter_rows(min_row=2):
        row[2].number_format = "0.0"
        row[3].number_format = "yyyy-mm-dd"
    for col, width in zip("ABCD", (14, 8, 12, 13)):
        laps.column_dimensions[col].width = width
    laps.freeze_panes = "A2"

    last = 1 + sum(len(t) for t in LAPS.values())
    names, times = f"{T['laps']}!$A$2:$A${last}", f"{T['laps']}!$C$2:$C${last}"
    res.append(T["results_header"])
    header(res, res[1])
    n = len(RUNNERS)
    for i, runner in enumerate(RUNNERS, start=2):
        res.append(
            [
                runner,
                f"=MINIFS({times},{names},A{i})",
                f"=AVERAGEIFS({times},{names},A{i})",
                f"=B{i}-MIN($B$2:$B${n + 1})",
                f"=RANK.EQ(B{i},$B$2:$B${n + 1},1)",
            ]
        )
        res[f"B{i}"].number_format = "0.0"
        res[f"C{i}"].number_format = "0.00"
        res[f"D{i}"].number_format = '+0.0;-0.0;"—"'
        res[f"E{i}"].alignment = Alignment(horizontal="center")

    res["G1"] = T["summary"]
    header(res, [res["G1"], res["H1"]])
    summary = [
        (T["fastest"], f"=MIN({times})", "0.0"),
        (T["by"], f"=XLOOKUP(MIN({times}),{times},{names})", None),
        (T["count"], f"=COUNTA({names})", "0"),
        (T["group_mean"], f"=AVERAGE({times})", "0.00"),
    ]
    for r, (label, formula, fmt) in enumerate(summary, start=2):
        res[f"G{r}"] = label
        res[f"H{r}"] = formula
        if fmt:
            res[f"H{r}"].number_format = fmt
    for col, width in zip("ABCDEFGH", (14, 14, 11, 17, 8, 3, 24, 12)):
        res.column_dimensions[col].width = width

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = T["chart_title"]
    chart.y_axis.title = T["seconds"]
    chart.y_axis.scaling.min = 50
    chart.y_axis.number_format = "0"
    chart.add_data(Reference(res, min_col=2, max_col=3, min_row=1, max_row=n + 1), titles_from_data=True)
    chart.set_categories(Reference(res, min_col=1, min_row=2, max_row=n + 1))
    chart.width, chart.height = 18, 9
    res.add_chart(chart, "A10")

    wb.active = 0
    wb.properties.creator = AUTHOR
    wb.properties.created = wb.properties.modified = WHEN
    with tempfile.TemporaryDirectory() as tmp:
        raw = Path(tmp) / "raw.xlsx"
        wb.save(raw)
        subprocess.run(
            ["cargo", "run", "-q", "-p", "oel-xlsx", "--example", "recalc", "--", str(raw), str(T["dir"] / T["xlsx"])],
            cwd=ROOT,
            check=True,
        )


# ── Deck ────────────────────────────────────────────────────────────────────────────────────


def text_props(frame, size, bold=False, color=None, align=None):
    for p in frame.paragraphs:
        if align is not None:
            p.alignment = align
        for r in p.runs:
            r.font.size = PPt(size)
            r.font.bold = bold
            if color:
                r.font.color.rgb = RGBColor.from_string(color)


def make_pptx() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    title_layout, content_layout, title_only, blank = (prs.slide_layouts[i] for i in (0, 1, 5, 6))

    # 1 · Intro with the logo.
    s = prs.slides.add_slide(title_layout)
    logo = Inches(1.9)
    s.shapes.add_picture(str(LOGO), (W - logo) // 2, Inches(0.9), logo, logo).name = "Logo"
    s.shapes.title.text = "OutilsEnLigne · Bureau"
    s.shapes.title.left, s.shapes.title.top, s.shapes.title.width, s.shapes.title.height = Inches(1), Inches(3.1), W - Inches(2), Inches(1.4)
    text_props(s.shapes.title.text_frame, 48, bold=True, color=INK, align=PP_ALIGN.CENTER)
    sub = s.placeholders[1]
    sub.text = T["deck_subtitle"]
    sub.left, sub.top, sub.width, sub.height = Inches(1), Inches(4.6), W - Inches(2), Inches(0.9)
    text_props(sub.text_frame, 24, color=VIOLET, align=PP_ALIGN.CENTER)

    # 2 · Three tips.
    s = prs.slides.add_slide(content_layout)
    s.shapes.title.text = T["how"]
    text_props(s.shapes.title.text_frame, 40, bold=True, color=INK)
    body = s.placeholders[1]
    body.left, body.top, body.width, body.height = Inches(0.9), Inches(1.9), W - Inches(1.8), Inches(4.6)
    tips = T["tips"]
    tf = body.text_frame
    tf.text = tips[0]
    for t in tips[1:]:
        tf.add_paragraph().text = t
    for p in tf.paragraphs:
        p.space_after = PPt(18)
    text_props(tf, 28, color=INK)

    # 3 · Race results: the workbook's numbers, replaced on the desk by a linked copy of its chart.
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = T["race"]
    text_props(s.shapes.title.text_frame, 40, bold=True, color=INK)
    data = CategoryChartData()
    data.categories = RUNNERS
    data.add_series(T["best"], [BEST[r] for r in RUNNERS])
    data.add_series(T["mean"], [MEAN[r] for r in RUNNERS])
    frame = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.2), Inches(1.8), W - Inches(2.4), Inches(5.0), data)
    frame.name = T["deck_chart"]
    chart = frame.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.minimum_scale = 50
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = T["seconds"]

    # 4 · Thank-you, white on black.
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor(0, 0, 0)
    box = s.shapes.add_textbox(Inches(1), (H - Inches(1.6)) // 2, W - Inches(2), Inches(1.6))
    box.name = T["thanks_name"]
    box.text_frame.text = T["thanks"]
    text_props(box.text_frame, 72, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)

    prs.core_properties.author = prs.core_properties.last_modified_by = AUTHOR
    prs.core_properties.title = T["deck_title"]
    prs.core_properties.created = prs.core_properties.modified = WHEN
    prs.save(T["dir"] / T["pptx"])


# ── Document ────────────────────────────────────────────────────────────────────────────────


def make_docx() -> None:
    doc = Document()
    doc.add_heading(T["doc_title"], level=0)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(LOGO), width=Cm(4))
    for text in T["doc_text"]:
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(8)
    doc.add_heading(T["shortcuts_title"], level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = T["shortcuts_header"]
    for what, keys in T["shortcuts"]:
        row = table.add_row().cells
        row[0].text, row[1].text = what, keys
    # Word ends the body with a paragraph, never a table.
    doc.add_paragraph().paragraph_format.space_before = Pt(8)
    doc.paragraphs[-1].add_run(T["doc_end"])

    doc.core_properties.author = doc.core_properties.last_modified_by = AUTHOR
    doc.core_properties.title = T["doc_title"]
    doc.core_properties.created = doc.core_properties.modified = WHEN
    doc.save(T["dir"] / T["docx"])


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--lang", choices=sorted(TEXT), default="fr")
    T = TEXT[parser.parse_args().lang]
    if not LOGO.exists():
        raise SystemExit(f"{LOGO} manque : voir l’en-tête de ce script (magick).")
    T["dir"].mkdir(exist_ok=True)
    make_xlsx()
    make_pptx()
    make_docx()
    for name in (T["xlsx"], T["pptx"], T["docx"]):
        f = T["dir"] / name
        print(f"{f.relative_to(ROOT)}  {f.stat().st_size:,} o")

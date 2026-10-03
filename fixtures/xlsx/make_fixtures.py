"""Builds the oel-xlsx test fixtures (run with a Python that has openpyxl).

rich.xlsx exercises what the codec models and what it must preserve; run
`soffice --headless --convert-to xlsx --outdir lo rich.xlsx` afterwards for the
LibreOffice-written twin (shared formulas, theme colours, calcChain).
"""
from datetime import date
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
ws = wb.active
ws.title = "Budget"
ws.sheet_properties.tabColor = "1F6F5C"
ws.append(["Poste", "Montant", "Part", "Date", "Code"])
rows = [("Loyer", 1450.5, date(2026, 1, 1), "007"), ("Épicerie", 612.25, date(2026, 1, 3), "012"),
        ("Transport", 180, date(2026, 1, 5), "100"), ("Café « bio »", 42.75, date(2026, 1, 9), "003")]
for i, (p, m, d, code) in enumerate(rows, start=2):
    ws.cell(i, 1, p); ws.cell(i, 2, m); ws.cell(i, 3, f"=B{i}/$B$7"); ws.cell(i, 4, d); ws.cell(i, 5, code)
ws["A7"] = "Total"; ws["B7"] = "=SUM(B2:B5)"; ws["C7"] = "=SUM(C2:C5)"
ws["A8"] = "Moyenne"; ws["B8"] = "=AVERAGE(B2:B5)"
ws["A9"] = "Max"; ws["B9"] = "=MAX(Montants)"
ws["A10"] = "Texte"; ws["B10"] = '=CONCATENATE(A2," : ",TEXT(B2,"0.00"))'
ws["A11"] = "Vrai"; ws["B11"] = True
ws["A12"] = "Erreur"; ws["B12"] = "=1/0"
ws["A13"] = "Multi\nligne"; ws["A13"].alignment = Alignment(wrap_text=True, vertical="top")
for c in "BC": ws[f"{c}7"].font = Font(bold=True)
for r in range(2, 10): ws[f"B{r}"].number_format = '#,##0.00 "$"'
for r in range(2, 8): ws[f"C{r}"].number_format = "0.0%"
for r in range(2, 6): ws[f"D{r}"].number_format = "yyyy-mm-dd"
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=12)
    c.fill = PatternFill("solid", fgColor="1F6F5C")
    c.alignment = Alignment(horizontal="center")
    c.border = Border(bottom=Side(style="medium", color="000000"))
ws["A7"].font = Font(italic=True, underline="single", strike=False)
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 14
ws.column_dimensions["E"].hidden = True
ws.row_dimensions[13].height = 32
ws.merge_cells("A15:C15"); ws["A15"] = "Cellules fusionnées"
ws.freeze_panes = "B2"
ws.auto_filter.ref = "A1:E5"
dv = DataValidation(type="list", formula1='"Oui,Non,Peut-être"', allow_blank=True); ws.add_data_validation(dv); dv.add("F2:F5")
ws.conditional_formatting.add("B2:B5", CellIsRule(operator="greaterThan", formula=["500"], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(color="9C0006")))
ws.conditional_formatting.add("C2:C5", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="1F6F5C"))
ws.conditional_formatting.add("B2:B5", DataBarRule(start_type="min", end_type="max", color="638EC6"))
ws.conditional_formatting.add("A2:A5", FormulaRule(formula=['LEN(A2)>8'], font=Font(bold=True)))
ws["A2"].hyperlink = "https://outilsenligne.ca"; ws["A2"].style = "Hyperlink"
ws["A3"].comment = Comment("Note de test", "JB")
wb.defined_names["Montants"] = DefinedName("Montants", attr_text="Budget!$B$2:$B$5")
chart = BarChart(); chart.title = "Montants"
chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
ws.add_chart(chart, "H2")

ws2 = wb.create_sheet("Données")
for r in range(1, 51):
    ws2.cell(r, 1, r); ws2.cell(r, 2, f"=A{r}*2"); ws2.cell(r, 3, f"=Budget!B$7+A{r}")
ws2["E1"] = "=SUMPRODUCT(A1:A50,B1:B50)"; ws2["E2"] = "=VLOOKUP(10,A1:B50,2,FALSE)"; ws2["E3"] = "=IFERROR(INDEX(B1:B50,MATCH(25,A1:A50,0)),\"?\")"
ws2["E4"] = '=COUNTIF(A1:A50,">25")'; ws2["E5"] = "=ROUND(PI(),4)"; ws2["E6"] = '=IF(E5>3,"grand","petit")'
ws3 = wb.create_sheet("Caché"); ws3["A1"] = "secret"; ws3.sheet_state = "hidden"
wb.save("rich.xlsx")
